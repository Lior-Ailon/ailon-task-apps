import json,sys,re,zipfile,xml.etree.ElementTree as ET
from openpyxl import load_workbook
path,variant=sys.argv[1:3]
w=load_workbook(path)
s=w.active
layout={'bakery':(27,46,47,38,39),'patisserie':(26,50,51,26,27),'ta':(10,37,38,10,11)}
first,conclusion,steps,water1,water2=layout[variant]
if variant=='bakery':
 first_gram=s['E9'].value
 second=s['E39'].value
elif variant=='patisserie':
 first_gram=s['E9'].value
 second=s['E26'].value
else:
 first_gram=s['C10'].value
 second=s['E12'].value
with zipfile.ZipFile(path) as archive:
 root=ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
 ns={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
 sorted_rows=all(
  [int.from_bytes(''.join(filter(str.isalpha,c.attrib.get('r',''))).encode(),'big') for c in row.findall('x:c',ns)] ==
  sorted(int.from_bytes(''.join(filter(str.isalpha,c.attrib.get('r',''))).encode(),'big') for c in row.findall('x:c',ns))
  for row in root.findall('.//x:row',ns))
result={
 'sortedCellOrder':sorted_rows,
 'lastSupplementalValue':next((s[f'Q{r}'].value for r in range(s.max_row,2,-1) if s[f'Q{r}'].value is not None),None),
 'sheets':len(w.worksheets),'images':len(s._images),
 'project':s['B4'].value,'researcher':s['B6'].value,
 'conclusions':s[f'C{conclusion}'].value,'nextSteps':s[f'C{steps}'].value,
 'firstGram':first_gram,'secondExperiment':second,
 'errors':sum(c.data_type=='e' for row in s for c in row),
 'creator':w.properties.creator,
 'printArea':str(s.print_area),
 'percentFormula':s[{'bakery':'F9','patisserie':'F9','ta':'D10'}[variant]].value,
 'percentCached':load_workbook(path,data_only=True).active[{'bakery':'F9','patisserie':'F9','ta':'D10'}[variant]].value,
 'merged':len(s.merged_cells.ranges),
}
print(json.dumps(result,ensure_ascii=False))
